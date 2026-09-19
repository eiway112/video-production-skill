#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
文档 → Markdown 转换器

支持格式:
  - PPTX (python-pptx): 按 slide 拆分章节，保留层级结构
  - DOCX (python-docx): 按段落/标题层级拆分，保留列表和表格

输出:
  - Markdown 文件 (UTF-8)
  - 提取的图片保存到 images/ 子目录

用法:
  python doc_to_markdown.py input.pptx
  python doc_to_markdown.py input.docx
  python doc_to_markdown.py input.pptx -o output.md
"""

import sys
import io
import argparse
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

# PPT 原生图片格式：原字节直接落盘（免二次编码损失）；其余（webp 等）转 PNG。
_NATIVE_IMAGE_EXTS = {"PNG": "png", "JPEG": "jpg", "GIF": "gif",
                      "BMP": "bmp", "TIFF": "tif"}


def _picture_blob(shape):
    """取 PICTURE shape 的原始字节，取不到返回 None。

    不能用 shape.image：python-pptx 只对 image_content_types 白名单内的部件返回
    ImagePart，webp 等未在 [Content_Types].xml 注册的部件退化为通用 Part，
    shape.image 会抛 AttributeError('Part' object has no attribute 'image')。
    """
    try:
        rId = shape._element.blip_rId
        if not rId:
            return None
        return shape.part.related_part(rId).blob
    except Exception:
        return None


def _write_picture(blob, img_dir, stem):
    """按真实格式写盘，返回文件名；不可解码的图片格式返回 None（交调用方记为跳过）。"""
    if not blob:
        return None
    from PIL import Image
    try:
        im = Image.open(io.BytesIO(blob))
        fmt = (im.format or "").upper()
    except Exception:
        return None
    if fmt in _NATIVE_IMAGE_EXTS:
        name = f"{stem}.{_NATIVE_IMAGE_EXTS[fmt]}"
        (img_dir / name).write_bytes(blob)
        return name
    name = f"{stem}.png"
    try:
        im.load()
        im.save(img_dir / name, "PNG")
    except Exception:
        return None
    return name


def convert_pptx(filepath: Path, output_dir: Path) -> str:
    """PPTX → Markdown: 每个 slide 一个章节"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    import base64

    prs = Presentation(str(filepath))
    img_dir = output_dir / "images"
    img_dir.mkdir(parents=True, exist_ok=True)

    lines = []
    lines.append(f"# {filepath.stem}\n")

    total_slides = len(prs.slides)
    img_count = 0
    skipped = []

    for i, slide in enumerate(prs.slides, 1):
        layout = slide.slide_layout.name if slide.slide_layout else "unknown"
        lines.append(f"\n## Slide {i} / {total_slides}\n")

        # Extract text content
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if not text:
                        continue
                    level = para.level
                    if level == 0:
                        # Check if it looks like a heading (short, no punctuation end)
                        if len(text) < 40 and not text[-1] in "。，、；：！？":
                            lines.append(f"**{text}**\n")
                        else:
                            lines.append(f"{text}\n")
                    else:
                        indent = "  " * (level - 1)
                        lines.append(f"{indent}- {text}\n")

            # Extract images
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img_count += 1
                stem = f"slide{i}_img{img_count}"
                name = _write_picture(_picture_blob(shape), img_dir, stem)
                if name:
                    lines.append(f"\n![{shape.name}](images/{name})\n")
                else:
                    skipped.append(f"{stem} ({shape.name})")

            # Extract tables
            if shape.has_table:
                table = shape.table
                lines.append("")
                # Header row
                headers = [cell.text.strip() for cell in table.rows[0].cells]
                lines.append("| " + " | ".join(headers) + " |")
                lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                # Data rows
                for row in list(table.rows)[1:]:
                    cells = [cell.text.strip() for cell in row.cells]
                    lines.append("| " + " | ".join(cells) + " |")
                lines.append("")

    if skipped:
        print(f"  跳过 {len(skipped)} 个无法解码的图片部件: {', '.join(skipped)}")

    return "\n".join(lines)


def convert_docx(filepath: Path, output_dir: Path) -> str:
    """DOCX → Markdown: 按标题层级拆分"""
    from docx import Document
    from docx.table import Table

    doc = Document(str(filepath))
    img_dir = output_dir / "images"
    img_dir.mkdir(parents=True, exist_ok=True)

    lines = []
    lines.append(f"# {filepath.stem}\n")

    img_count = 0

    for element in doc.element.body:
        tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

        if tag == "p":
            # Paragraph
            from docx.text.paragraph import Paragraph
            para = Paragraph(element, doc)
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name if para.style else ""

            # Heading levels
            if "Heading 1" in style_name or style_name == "heading 1":
                lines.append(f"\n## {text}\n")
            elif "Heading 2" in style_name or style_name == "heading 2":
                lines.append(f"\n### {text}\n")
            elif "Heading 3" in style_name or style_name == "heading 3":
                lines.append(f"\n#### {text}\n")
            elif "Title" in style_name:
                lines.append(f"\n# {text}\n")
            elif "List" in style_name:
                lines.append(f"- {text}\n")
            else:
                lines.append(f"{text}\n")

        elif tag == "tbl":
            # Table
            table = Table(element, doc)
            lines.append("")
            for row_idx, row in enumerate(table.rows):
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                lines.append("| " + " | ".join(cells) + " |")
                if row_idx == 0:
                    lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
            lines.append("")

    # Extract embedded images from relationships
    skipped = []
    for rel_id, rel in doc.part.rels.items():
        if "image" in rel.reltype:
            img_count += 1
            name = _write_picture(rel.target_part.blob, img_dir, f"doc_img{img_count}")
            if name:
                lines.append(f"\n![image](images/{name})\n")
            else:
                skipped.append(f"doc_img{img_count}")

    if skipped:
        print(f"  跳过 {len(skipped)} 个无法解码的图片部件: {', '.join(skipped)}")

    return "\n".join(lines)


def convert(filepath: Path, output_path: Path = None) -> Path:
    """Auto-detect format and convert to Markdown"""
    suffix = filepath.suffix.lower()
    output_dir = filepath.parent

    if output_path is None:
        output_path = filepath.with_suffix(".md")
    else:
        output_dir = output_path.parent

    output_dir.mkdir(parents=True, exist_ok=True)

    converters = {
        ".pptx": convert_pptx,
        ".docx": convert_docx,
    }

    converter = converters.get(suffix)
    if not converter:
        supported = ", ".join(converters.keys())
        raise ValueError(f"Unsupported format: {suffix} (supported: {supported})")

    print(f"Converting: {filepath}")
    print(f"Format: {suffix}")

    md_content = converter(filepath, output_dir)

    output_path.write_text(md_content, encoding="utf-8")

    # Stats
    line_count = len(md_content.split("\n"))
    img_count = md_content.count("![")
    table_count = md_content.count("| ---")

    print(f"Output: {output_path}")
    print(f"  Lines: {line_count}")
    print(f"  Images: {img_count}")
    print(f"  Tables: {table_count}")

    return output_path


def main():
    parser = argparse.ArgumentParser(description="Document to Markdown converter")
    parser.add_argument("input", type=Path, help="Input file (PPTX or DOCX)")
    parser.add_argument("-o", "--output", type=Path, help="Output Markdown path")
    args = parser.parse_args()

    if not args.input.exists():
        print(f"ERROR: file not found: {args.input}")
        sys.exit(1)

    convert(args.input, args.output)


if __name__ == "__main__":
    main()
